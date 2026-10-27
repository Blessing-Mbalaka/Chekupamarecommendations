from unittest.mock import patch

from django.test import TestCase

from ingestion.models import ContentChunk
from ingestion.services.vector_store import CHUNKING_STRATEGY, index_material, retrieve_chunks, split_document
from learning.models import Course, Material


class SemanticChunkingTests(TestCase):
    def setUp(self):
        self.course = Course.objects.create(code="RAG101", title="Retrieval")
        self.material = Material.objects.create(
            course=self.course,
            title="Systems handbook",
            source_type=Material.SourceType.FILE,
            is_validated=True,
        )

    def test_split_document_retains_structure_and_uses_bounded_chunks(self):
        first_page = "INTRODUCTION\n\n" + " ".join(
            f"Feedback loops sentence number {index} explains reinforcing system behaviour."
            for index in range(45)
        )
        second_page = "APPLICATIONS\n\n" + " ".join(
            f"A practical boundary example number {index} demonstrates applied systems analysis."
            for index in range(45)
        )

        chunks = split_document(f"{first_page}\f{second_page}")

        self.assertGreaterEqual(len(chunks), 2)
        self.assertEqual(chunks[0].page_number, 1)
        self.assertEqual(chunks[0].section_title, "INTRODUCTION")
        self.assertTrue(any(chunk.page_number == 2 for chunk in chunks))
        self.assertTrue(all(len(chunk.text.split()) <= 510 for chunk in chunks))

    @patch("ingestion.services.vector_store.embed_text", return_value=([0.2, 0.4], "test"))
    def test_index_material_persists_rag_provenance(self, mock_embed):
        chunks = index_material(
            self.material,
            text="Feedback loops connect actions and consequences across a system. "
            "A balancing loop counters change and supports stability.",
        )

        self.assertEqual(len(chunks), 1)
        chunk = ContentChunk.objects.get(material=self.material)
        self.assertEqual(chunk.chunking_strategy, CHUNKING_STRATEGY)
        self.assertEqual(chunk.embedding_backend, "test")
        self.assertEqual(len(chunk.content_hash), 64)
        self.assertGreater(chunk.token_count, 0)
        self.assertEqual(chunk.metadata["course_id"], self.course.pk)
        mock_embed.assert_called_once()

    @patch("ingestion.services.vector_store.embed_texts", return_value=[([0.1], "test"), ([0.2], "test")])
    def test_pptx_is_extracted_per_slide_with_slide_numbers(self, _mock_embed):
        import io

        from django.core.files.base import ContentFile
        from pptx import Presentation

        deck = Presentation()
        for title, body in (("Feedback loops", "Reinforcing loops amplify change. " * 40),
                            ("Boundaries", "A system boundary defines scope. " * 40)):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        buffer = io.BytesIO()
        deck.save(buffer)
        self.material.file.save("deck.pptx", ContentFile(buffer.getvalue()), save=True)

        chunks = index_material(self.material)

        self.assertTrue(all(chunk.page_number is not None for chunk in chunks))
        self.assertIn("Reinforcing loops amplify change.", " ".join(chunk.text for chunk in chunks))
        self.assertIn("system boundary", " ".join(chunk.text for chunk in chunks))


class HybridRetrievalTests(TestCase):
    def setUp(self):
        self.course = Course.objects.create(code="SEARCH101", title="Search")
        self.material = Material.objects.create(
            course=self.course,
            title="Control systems",
            is_validated=True,
        )

    @patch("ingestion.services.vector_store.embed_text", return_value=([], "fallback"))
    def test_bm25_ranks_exact_course_term_without_embeddings(self, mock_embed):
        ContentChunk.objects.create(
            material=self.material,
            ordinal=0,
            text="Homeostasis is maintained through a balancing feedback loop.",
        )
        ContentChunk.objects.create(
            material=self.material,
            ordinal=1,
            text="A broad overview of organizations and management practice.",
        )

        matches = retrieve_chunks("Explain homeostasis", course=self.course, limit=2)

        self.assertEqual(matches[0]["chunk"].ordinal, 0)
        self.assertGreater(matches[0]["bm25_score"], 0)
        self.assertEqual(matches[0]["semantic_score"], 0)
        mock_embed.assert_called_once_with("Explain homeostasis")

    @patch("ingestion.services.vector_store.embed_text", return_value=([1.0, 0.0], "test"))
    def test_hybrid_score_combines_semantic_and_bm25_signals(self, mock_embed):
        ContentChunk.objects.create(
            material=self.material,
            ordinal=0,
            text="Leverage points can produce disproportionate system change.",
            embedding=[1.0, 0.0],
        )
        ContentChunk.objects.create(
            material=self.material,
            ordinal=1,
            text="A generic unrelated passage.",
            embedding=[0.0, 1.0],
        )

        matches = retrieve_chunks("system leverage points", course=self.course, limit=2)

        self.assertEqual(matches[0]["chunk"].ordinal, 0)
        self.assertGreater(matches[0]["score"], 0)
        self.assertGreater(matches[0]["semantic_score"], 0)
        self.assertGreater(matches[0]["bm25_score"], 0)

    def test_retrieval_rejects_missing_course_scope(self):
        with self.assertRaisesRegex(ValueError, "course is required"):
            retrieve_chunks("anything", course=None)

    @patch("ingestion.services.vector_store.embed_text", return_value=([], "fallback"))
    def test_retrieval_never_returns_another_courses_chunks(self, _mock_embed):
        other_course = Course.objects.create(code="OTHER101", title="Other Course")
        other_material = Material.objects.create(course=other_course, title="Other notes", is_validated=True)
        ContentChunk.objects.create(material=self.material, ordinal=0, text="sharedterm allowed course evidence")
        ContentChunk.objects.create(material=other_material, ordinal=0, text="sharedterm forbidden course evidence")

        matches = retrieve_chunks("sharedterm", course=self.course, limit=10)

        self.assertTrue(matches)
        self.assertEqual({match["material"].course_id for match in matches}, {self.course.pk})
